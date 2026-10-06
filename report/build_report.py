"""Build reports and a secret-free submission archive. Run from any directory.
Requires: python -m pip install reportlab python-docx Pillow
"""
from pathlib import Path
from xml.sax.saxutils import escape
import base64
import html
import re
import zipfile
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak, Preformatted
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib import colors
from docx import Document
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
styles = getSampleStyleSheet()
styles.add(ParagraphStyle('CodeSmall', fontName='Courier', fontSize=6.5, leading=8))
story = []
doc = Document()
doc.core_properties.author = 'Matvei Krautsou'
doc.core_properties.title = 'Lab 3 — Campus Shop'
parts = ['<!doctype html><meta charset="utf-8"><title>Matvei Krautsou — Lab 3</title><style>body{font:16px Arial;max-width:900px;margin:40px auto;padding:20px;line-height:1.5}img{display:block;max-width:100%;max-height:760px;margin:auto}pre{font:12px monospace;white-space:pre-wrap;background:#f1f4f8;padding:16px}h1,h2,h3{color:#193c6a}@media print{img{max-height:650px}h2{break-after:avoid}}</style>']
for line in (HERE / 'REPORT.md').read_text().splitlines():
    if line.startswith('!['):
        name = re.search(r'\]\((.*?)\)', line).group(1)
        path = HERE / name
        story.append(Image(str(path), width=260, height=520, kind='proportional'))
        doc.add_picture(str(path), width=Inches(3.2))
        data = base64.b64encode(path.read_bytes()).decode()
        parts.append(f'<img src="data:image/png;base64,{data}">')
    elif line.startswith('#'):
        level = len(line) - len(line.lstrip('#'))
        text = line.lstrip('# ')
        story.append(Paragraph(escape(text), styles['Title' if level == 1 else 'Heading2']))
        doc.add_heading(text, level=min(level, 3))
        n = min(level, 3)
        parts.append(f'<h{n}>{html.escape(text)}</h{n}>')
    elif line:
        story.append(Paragraph(escape(line), styles['BodyText']))
        doc.add_paragraph(line)
        parts.append(f'<p>{html.escape(line)}</p>')
    else:
        story.append(Spacer(1, 8))

sources = ['backend/package.json', 'backend/server.js', 'backend/index.html',
           'backend/server.test.js', 'bot/main.py', 'bot/test_bot.py',
           '.env.example', '.gitignore', 'README.md', 'report/build_report.py']
for name in sources:
    text = (ROOT / name).read_text()
    if re.search(r'\d{8,12}:[A-Za-z0-9_-]{30,}', text):
        raise RuntimeError('Token-like content detected in ' + name)
    story.extend([PageBreak(), Paragraph(escape(name), styles['Heading2'])])
    # Preserve every source character, wrapping only the printed representation.
    import textwrap
    printed = '\n'.join(piece for line in text.expandtabs(4).splitlines()
                        for piece in (textwrap.wrap(line, width=105, replace_whitespace=False,
                                                  drop_whitespace=False) or ['']))
    story.append(Preformatted(printed, styles['CodeSmall']))
    doc.add_page_break()
    doc.add_heading(name, level=2)
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = 'Courier New'
    run.font.size = Pt(7)
    parts.append(f'<h2>{html.escape(name)}</h2><pre>{html.escape(text)}</pre>')

def footer(canvas, document):
    canvas.setFont('Helvetica', 8)
    canvas.setFillColor(colors.HexColor('#667085'))
    canvas.drawString(42, 24, 'Matvei Krautsou | RPA Lab 3 | Campus Shop')
    canvas.drawRightString(553, 24, str(document.page))

SimpleDocTemplate(str(HERE / 'Lab3-report.pdf'), title='Lab 3 — Campus Shop',
                  author='Matvei Krautsou', leftMargin=42, rightMargin=42,
                  topMargin=42, bottomMargin=42).build(story, onFirstPage=footer, onLaterPages=footer)
doc.save(HERE / 'Lab3-report.docx')
(HERE / 'Lab3-report.html').write_text('\n'.join(parts))
archive = ROOT / 'Lab3-Matvei-Krautsou.zip'
with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as z:
    for path in sorted(ROOT.rglob('*')):
        relative = path.relative_to(ROOT)
        if (not path.is_file() or path == archive or path.suffix == '.zip'
                or path.name == '.env' or any(p in {'.git', '__pycache__', '.venv', 'data'} for p in relative.parts)):
            continue
        z.write(path, Path('lab3') / relative)
print('Generated PDF, DOCX, standalone HTML and secret-free submission ZIP')
